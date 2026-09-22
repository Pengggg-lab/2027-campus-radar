#!/usr/bin/env python3
"""Build the public company dataset from the source Excel workbook."""
from __future__ import annotations

import argparse
import json
from datetime import datetime, timezone
from pathlib import Path
from typing import Any

import openpyxl


SHEET_META = {
    "央企": {"group": "央企", "region": "全国"},
    "四川国企": {"group": "地方国企", "region": "四川"},
    "重庆国企": {"group": "地方国企", "region": "重庆"},
    "贵州国企": {"group": "地方国企", "region": "贵州"},
}


def value(cell: Any) -> Any:
    cell_value = cell.value
    return None if cell_value is None else str(cell_value).strip()


def build(source: Path, output: Path, previous: Path | None = None, overrides: Path | None = None) -> dict[str, Any]:
    wb = openpyxl.load_workbook(source, data_only=False)
    override_data: dict[str, dict[str, Any]] = {}
    if overrides and overrides.exists():
        override_data = json.loads(overrides.read_text(encoding="utf-8"))
    old_by_id: dict[str, dict[str, Any]] = {}
    old_by_name: dict[str, dict[str, Any]] = {}
    old_meta: dict[str, Any] = {}
    if previous and previous.exists():
        old_data = json.loads(previous.read_text(encoding="utf-8"))
        old_meta = old_data.get("meta", {})
        old_by_id = {item["id"]: item for item in old_data.get("companies", [])}
        old_by_name = {item["name"]: item for item in old_data.get("companies", [])}

    companies: list[dict[str, Any]] = []
    sequence = 0
    for ws in wb.worksheets:
        if ws.title not in SHEET_META:
            continue
        meta = SHEET_META[ws.title]
        for row in range(3, ws.max_row + 1):
            name_cell = ws.cell(row, 2)
            name = value(name_cell)
            if not name:
                continue
            sequence += 1
            company_id = f"c{sequence:04d}"
            official_cell = ws.cell(row, 3)
            campus_cell = ws.cell(row, 4)
            official_url = official_cell.hyperlink.target if official_cell.hyperlink else None
            campus_url = campus_cell.hyperlink.target if campus_cell.hyperlink else None
            override = override_data.get(name, {})
            official_url = override.get("officialUrl") or official_url
            campus_url = override.get("campusUrl") or campus_url
            old = old_by_name.get(name) or old_by_id.get(company_id, {})
            item = {
                "id": company_id,
                "name": name,
                "group": meta["group"],
                "region": meta["region"],
                "sheet": ws.title,
                "sourceRow": row,
                "officialUrl": official_url,
                "campusUrl": campus_url,
                "status": old.get("status", "not_checked"),
                "statusLabel": old.get("statusLabel", "待扫描"),
                "todayUpdate": bool(old.get("todayUpdate", False)),
                "evidence": old.get("evidence"),
                "evidenceDate": old.get("evidenceDate"),
                "pageTitle": old.get("pageTitle"),
                "finalUrl": old.get("finalUrl"),
                "relatedLinks": old.get("relatedLinks", []),
                "subsidiaryHint": bool(old.get("subsidiaryHint", False)),
                "lastChecked": old.get("lastChecked"),
                "checkError": old.get("checkError"),
                "previousStatus": old.get("previousStatus"),
            }
            companies.append(item)

    meta = {
        "title": "2027届秋招雷达",
            "targetYear": 2027,
            "generatedAt": datetime.now(timezone.utc).isoformat(),
            "companyCount": len(companies),
            "source": source.name,
        }
    for key in ("scannedAt", "stats", "browserChecks"):
        if key in old_meta:
            meta[key] = old_meta[key]
    payload = {"meta": meta, "companies": companies}
    output.parent.mkdir(parents=True, exist_ok=True)
    output.write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")
    return payload


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--source", type=Path, default=Path("data/source.xlsx"))
    parser.add_argument("--output", type=Path, default=Path("site/data/companies.json"))
    parser.add_argument("--previous", type=Path)
    parser.add_argument("--overrides", type=Path, default=Path("data/overrides.json"))
    args = parser.parse_args()
    previous = args.previous or args.output
    payload = build(args.source, args.output, previous, args.overrides)
    print(f"built {payload['meta']['companyCount']} companies -> {args.output}")


if __name__ == "__main__":
    main()
